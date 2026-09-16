"""Bounded document extraction with production fail-closed malware scanning."""
import io
import re
import socket
from collections import Counter
from copy import deepcopy
from dataclasses import dataclass, field
from html import escape
from html.parser import HTMLParser
from typing import ClassVar
from urllib.parse import urlparse

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph as DocxParagraph
from fastapi import HTTPException, UploadFile
from pypdf import PdfReader
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .config import Settings

SUPPORTED_DOCUMENT_TYPES = {
    "text/plain": "text",
    "text/markdown": "text",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
}

SECTION_HEADINGS = {
    "SUMMARY", "PROFESSIONAL SUMMARY", "CAREER SUMMARY", "PROFILE", "PROFESSIONAL PROFILE", "CAREER PROFILE", "OBJECTIVE", "SKILLS", "CORE SKILLS",
    "TECHNICAL SKILLS", "TECHNICAL PROFICIENCIES", "EXPERIENCE", "PROFESSIONAL EXPERIENCE", "RELEVANT EXPERIENCE", "SELECTED EXPERIENCE",
    "WORK EXPERIENCE", "WORK HISTORY", "EMPLOYMENT", "EMPLOYMENT HISTORY", "CAREER HISTORY", "PROFESSIONAL BACKGROUND",
    "PROJECTS", "SELECTED PROJECTS", "SELECTED PRODUCTS", "ACADEMIC PROJECTS", "EDUCATION", "CERTIFICATIONS", "AWARDS", "VOLUNTEERING",
}

URL_PATTERN = re.compile(r"(?:https?://|www\.|mailto:)[^\s<>\"']+", re.IGNORECASE)
URL_TRAILING_PUNCTUATION = ".,;:!?"


def is_section_heading(line: str) -> bool:
    normalized = " ".join(line.replace(":", "").split()).upper()
    return normalized in SECTION_HEADINGS


def paragraph_text(document: Document) -> str:
    """Keep paragraph breaks and table rows so layout meaning survives extraction."""
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [" ".join(cell.text.split()) for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n\n".join(parts)


def pdf_text(reader: PdfReader) -> str:
    """Prefer pypdf layout mode so ordinary PDF text does not become fragments."""
    pages: list[str] = []
    for page in reader.pages:
        try:
            extracted = page.extract_text(extraction_mode="layout")
        except TypeError:
            extracted = page.extract_text()
        if extracted:
            pages.append(extracted.strip())
    return "\n\n".join(pages)


ALLOWED_FONT_FAMILIES = {"Aptos", "Arial", "Calibri", "Georgia", "Times New Roman"}


def _point_size(value) -> float | None:
    return float(value.pt) if value is not None else None


def docx_style_profile(document: Document) -> dict[str, str | float | bool]:
    """Extract portable DOCX typography signals without retaining source document XML."""
    paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
    normal = document.styles["Normal"]
    font_names = [run.font.name for paragraph in paragraphs for run in paragraph.runs if run.font.name in ALLOWED_FONT_FAMILIES]
    font_family = Counter(font_names).most_common(1)[0][0] if font_names else normal.font.name if normal.font.name in ALLOWED_FONT_FAMILIES else "Aptos"
    body_paragraphs = [paragraph for index, paragraph in enumerate(paragraphs) if index > 1 and not is_section_heading(paragraph.text)]
    body_sizes = [size for paragraph in body_paragraphs for run in paragraph.runs if (size := _point_size(run.font.size)) is not None]
    body_size = min(14, max(8, sum(body_sizes) / len(body_sizes))) if body_sizes else 10.5
    heading_paragraphs = [paragraph for paragraph in paragraphs if is_section_heading(paragraph.text)]
    heading_sizes = [size for paragraph in heading_paragraphs for run in paragraph.runs if (size := _point_size(run.font.size)) is not None]
    heading_size = min(16, max(9, sum(heading_sizes) / len(heading_sizes))) if heading_sizes else max(10.5, body_size)
    name_sizes = [size for run in paragraphs[0].runs if (size := _point_size(run.font.size)) is not None] if paragraphs else []
    name_size = min(28, max(12, max(name_sizes))) if name_sizes else max(18, heading_size + 5)
    emphasis_candidates = [paragraph for paragraph in body_paragraphs if not paragraph.text.lstrip().startswith(("-", "*", "•"))]
    emphasize_role_lines = any(sum(len(run.text) for run in paragraph.runs if run.bold) >= len(paragraph.text.strip()) * 0.6 for paragraph in emphasis_candidates if paragraph.text.strip())
    italic_metadata = any(bool(run.italic) and re.search(r"\b(?:19|20)\d{2}\b", paragraph.text) for paragraph in paragraphs for run in paragraph.runs)
    return {
        "font_family": font_family,
        "body_size": round(body_size, 1),
        "line_height": round(min(20, max(10, body_size * 1.25)), 1),
        "name_size": round(name_size, 1),
        "heading_size": round(heading_size, 1),
        "heading_uppercase": all(paragraph.text.strip() == paragraph.text.strip().upper() for paragraph in heading_paragraphs) if heading_paragraphs else True,
        "emphasize_role_lines": emphasize_role_lines,
        "italic_metadata": italic_metadata,
    }


def docx_entry_lines(document: Document) -> list[str]:
    """Return bold DOCX entry paragraphs so the editor can preserve semantic boundaries."""
    entries: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        bold_characters = sum(len(run.text) for run in paragraph.runs if run.bold)
        if text and not is_section_heading(text) and bold_characters >= len(text) * 0.6:
            entries.append(text)
    return entries[:500]


def docx_editor_html(document: Document) -> str:
    """Create safe editable HTML from DOCX paragraphs and styled runs."""
    blocks: list[str] = []
    open_list = False
    for paragraph in document.paragraphs:
        if not paragraph.text.strip():
            continue
        is_bullet = "List Bullet" in (paragraph.style.name or "") or paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None
        if open_list and not is_bullet:
            blocks.append("</ul>")
            open_list = False
        rendered_runs: list[str] = []
        for run in paragraph.runs:
            value = escape(run.text).replace("\n", "<br/>")
            for match in reversed(list(URL_PATTERN.finditer(run.text))):
                visible = match.group(0).rstrip(URL_TRAILING_PUNCTUATION)
                href = f"https://{visible}" if visible.lower().startswith("www.") else visible
                if visible and safe_link(href):
                    start, end = match.span()
                    value = f'{escape(run.text[:start])}<a href="{escape(href, quote=True)}" rel="noreferrer" target="_blank">{escape(visible)}</a>{escape(run.text[end:])}'
                    break
            if run.bold:
                value = f"<strong>{value}</strong>"
            if run.italic:
                value = f"<em>{value}</em>"
            if run.underline:
                value = f"<u>{value}</u>"
            rendered_runs.append(value)
        content = "".join(rendered_runs) or escape(paragraph.text)
        alignment = {WD_ALIGN_PARAGRAPH.CENTER: "center", WD_ALIGN_PARAGRAPH.RIGHT: "right"}.get(paragraph.alignment, "left")
        style = f' style="text-align: {alignment}"' if alignment != "left" else ""
        if is_bullet:
            if not open_list:
                blocks.append("<ul>")
                open_list = True
            blocks.append(f"<li{style}>{content}</li>")
        else:
            tag = "h3" if is_section_heading(paragraph.text) else "p"
            blocks.append(f"<{tag}{style}>{content}</{tag}>")
    if open_list:
        blocks.append("</ul>")
    return "".join(blocks)


@dataclass(frozen=True)
class ExtractedDocument:
    text: str
    page_count: int | None = None
    style_profile: dict[str, str | float | bool] | None = None
    entry_lines: list[str] | None = None
    editor_html: str | None = None


@dataclass(frozen=True)
class TextRun:
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    href: str | None = None


@dataclass
class ResumeBlock:
    kind: str
    runs: list[TextRun] = field(default_factory=list)
    alignment: str = "left"


def safe_link(value: str | None) -> str | None:
    if not value:
        return None
    parsed = urlparse(value)
    return value if parsed.scheme in {"http", "https", "mailto"} else None


def _linkify_run(run: TextRun) -> list[TextRun]:
    """Turn displayed, safe URLs into links while retaining all run styling."""
    if run.href:
        return [run]
    linked_runs: list[TextRun] = []
    cursor = 0
    for match in URL_PATTERN.finditer(run.text):
        matched = match.group(0)
        visible_url = matched.rstrip(URL_TRAILING_PUNCTUATION)
        if not visible_url:
            continue
        if match.start() > cursor:
            linked_runs.append(
                TextRun(run.text[cursor:match.start()], run.bold, run.italic, run.underline)
            )
        href = f"https://{visible_url}" if visible_url.lower().startswith("www.") else visible_url
        linked_runs.append(
            TextRun(visible_url, run.bold, run.italic, run.underline, safe_link(href))
        )
        cursor = match.start() + len(visible_url)
    if not linked_runs:
        return [run]
    if cursor < len(run.text):
        linked_runs.append(TextRun(run.text[cursor:], run.bold, run.italic, run.underline))
    return linked_runs


def _linkify_blocks(blocks: list[ResumeBlock]) -> list[ResumeBlock]:
    for block in blocks:
        block.runs = [linked_run for run in block.runs for linked_run in _linkify_run(run)]
    return blocks


class ResumeHtmlParser(HTMLParser):
    """Parse the editor's allow-listed formatting into a portable document model."""

    _block_tags: ClassVar[set[str]] = {"p", "div", "li", "h1", "h2", "h3"}

    def __init__(self) -> None:
        super().__init__()
        self._blocks: list[ResumeBlock] = []
        self._current: ResumeBlock | None = None
        self._styles: list[tuple[str, str | None]] = []

    def _flush(self) -> None:
        if self._current and any(run.text.strip() for run in self._current.runs):
            self._blocks.append(self._current)
        self._current = None

    def _start_block(self, kind: str, attrs: dict[str, str | None]) -> None:
        self._flush()
        style = attrs.get("style") or ""
        alignment = "center" if "text-align: center" in style else "right" if "text-align: right" in style else "left"
        self._current = ResumeBlock(kind=kind, alignment=alignment)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        attributes = dict(attrs)
        if tag in self._block_tags:
            self._start_block(tag, attributes)
        elif tag == "br":
            self._append("\n")
        elif tag in {"b", "strong", "i", "em", "u", "a"}:
            self._styles.append((tag, safe_link(attributes.get("href")) if tag == "a" else None))

    def handle_endtag(self, tag: str) -> None:
        if tag in self._block_tags:
            self._flush()
        elif tag in {"b", "strong", "i", "em", "u", "a"}:
            for index in range(len(self._styles) - 1, -1, -1):
                if self._styles[index][0] == tag:
                    self._styles.pop(index)
                    break

    def _append(self, text: str) -> None:
        if not text:
            return
        if self._current is None:
            self._current = ResumeBlock(kind="p")
        tags = {tag for tag, _ in self._styles}
        href = next((value for tag, value in reversed(self._styles) if tag == "a"), None)
        run = TextRun(text=text, bold=bool(tags & {"b", "strong"}), italic=bool(tags & {"i", "em"}), underline="u" in tags, href=href)
        if self._current.runs and self._current.runs[-1].bold == run.bold and self._current.runs[-1].italic == run.italic and self._current.runs[-1].underline == run.underline and self._current.runs[-1].href == run.href:
            prior = self._current.runs[-1]
            self._current.runs[-1] = TextRun(prior.text + run.text, prior.bold, prior.italic, prior.underline, prior.href)
        else:
            self._current.runs.append(run)

    def handle_data(self, data: str) -> None:
        self._append(data)

    def blocks(self) -> list[ResumeBlock]:
        self._flush()
        return self._blocks

    def text(self) -> str:
        rows = []
        for block in self.blocks():
            text = "".join(run.text for run in block.runs).strip()
            if text:
                rows.append(f"- {text}" if block.kind == "li" else text)
        return "\n".join(rows).strip()


def editor_html_to_text(resume_html: str | None, fallback: str) -> str:
    if not resume_html:
        return fallback
    parser = ResumeHtmlParser()
    parser.feed(resume_html)
    return parser.text() or fallback


def editor_html_matches_resume_text(resume_html: str | None, resume_text: str) -> bool:
    """Use rich editor markup only when it still represents the requested draft.

    The browser keeps plain text and HTML separately so a document can retain
    inline formatting.  A stale HTML snapshot must never replace the current
    tailored text during an export.
    """
    if not resume_html:
        return False

    def normalized_lines(value: str) -> list[str]:
        return [
            re.sub(r"^(?:[-*\u2022]\s+)", "", " ".join(line.split()))
            for line in value.splitlines()
            if line.strip()
        ]

    return normalized_lines(editor_html_to_text(resume_html, "")) == normalized_lines(
        resume_text
    )


def document_blocks(resume_text: str, resume_html: str | None) -> list[ResumeBlock]:
    if resume_html:
        parser = ResumeHtmlParser()
        parser.feed(resume_html)
        blocks = parser.blocks()
        if blocks:
            return _linkify_blocks(blocks)
    lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
    blocks = [ResumeBlock(kind="h1" if index == 0 else "p", runs=[TextRun(line[2:].strip() if line.startswith(("- ", "* ", "â€¢ ")) else line)], alignment="center" if index < 2 else "left") if not line.startswith(("- ", "* ", "â€¢ ")) else ResumeBlock(kind="li", runs=[TextRun(line[2:].strip())]) for index, line in enumerate(lines)]
    return _linkify_blocks(blocks)


class DocumentService:
    def __init__(self, settings: Settings) -> None:
        self._settings = settings

    async def extract(self, file: UploadFile) -> ExtractedDocument:
        document_type = SUPPORTED_DOCUMENT_TYPES.get(file.content_type or "")
        if not document_type:
            raise HTTPException(status_code=422, detail="Use a .txt, .md, .pdf, or .docx document.")
        data = await file.read(self._settings.max_import_bytes + 1)
        if len(data) > self._settings.max_import_bytes:
            raise HTTPException(status_code=413, detail="File exceeds the configured upload limit.")
        self._scan(data)
        try:
            style_profile = None
            entry_lines = None
            editor_html = None
            if document_type == "text":
                text, page_count = data.decode("utf-8", errors="replace"), None
            elif document_type == "pdf":
                reader = PdfReader(io.BytesIO(data))
                text, page_count = pdf_text(reader), len(reader.pages)
            else:
                document = Document(io.BytesIO(data))
                text, page_count, style_profile, entry_lines, editor_html = paragraph_text(document), None, docx_style_profile(document), docx_entry_lines(document), docx_editor_html(document)
        except Exception as error:
            raise HTTPException(status_code=422, detail="That document could not be read.") from error
        text = text.strip()
        if len(text) < 50: raise HTTPException(status_code=422, detail="The document is too short or has no readable text.")
        return ExtractedDocument(text=text[:100_000], page_count=page_count, style_profile=style_profile, entry_lines=entry_lines, editor_html=editor_html)

    def validate_docx_source(self, data: bytes) -> None:
        """Scan and open a transient DOCX before source-preserving export."""
        self._scan(data)
        try:
            Document(io.BytesIO(data))
        except Exception as error:
            raise HTTPException(status_code=422, detail="The original DOCX could not be opened.") from error

    def _scan(self, data: bytes) -> None:
        if not self._settings.clamav_host:
            if self._settings.environment == "development": return
            raise HTTPException(status_code=503, detail="Document scanning is not configured.")
        try:
            with socket.create_connection((self._settings.clamav_host, self._settings.clamav_port), timeout=5) as client:
                client.sendall(b"zINSTREAM\0")
                client.sendall(len(data).to_bytes(4, "big") + data + (0).to_bytes(4, "big"))
                response = client.recv(4096).decode("utf-8", errors="replace")
        except OSError as error:
            raise HTTPException(status_code=503, detail="Document scanning is unavailable.") from error
        if "OK" not in response or "FOUND" in response:
            raise HTTPException(status_code=422, detail="The document did not pass the malware scan.")


class ResumeExportService:
    """Render Rezzie's structured text contract as an editable ATS-friendly DOCX."""

    def render_docx(self, resume_text: str, target_page_count: int | None = None) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.65)
        section.left_margin = section.right_margin = Inches(0.7)

        normal = document.styles["Normal"]
        normal.font.name = "Aptos"
        compact = target_page_count == 1
        normal.font.size = Pt(10 if compact else 10.5)
        normal.paragraph_format.space_after = Pt(2 if compact else 4)

        lines = [line.strip() for line in resume_text.splitlines()]
        content_indices = [index for index, line in enumerate(lines) if line]
        name_index = content_indices[0] if content_indices else -1
        contact_index = content_indices[1] if len(content_indices) > 1 else -1

        for index, line in enumerate(lines):
            if not line:
                continue
            if index == name_index:
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = paragraph.add_run(line)
                run.bold = True
                run.font.size = Pt(18)
            elif index == contact_index:
                paragraph = document.add_paragraph(line)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.runs[0].font.size = Pt(9)
            elif is_section_heading(line):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(7 if compact else 10)
                paragraph.paragraph_format.space_after = Pt(3)
                run = paragraph.add_run(line.upper())
                run.bold = True
                run.font.size = Pt(11)
            elif line.startswith(("- ", "* ", "• ")):
                document.add_paragraph(line[2:].strip(), style="List Bullet")
            else:
                document.add_paragraph(line)

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    def render_pdf(self, resume_text: str, target_page_count: int | None = None) -> bytes:
        """Fit to the source page target when practical; never truncate resume content."""
        scales = (1.0, 0.95, 0.9, 0.85) if target_page_count == 1 else (1.0,)
        rendered = b""
        for scale in scales:
            rendered = self._render_pdf(resume_text, scale)
            if target_page_count is None or len(PdfReader(io.BytesIO(rendered)).pages) <= target_page_count:
                return rendered
        return rendered

    @staticmethod
    def _render_pdf(resume_text: str, scale: float) -> bytes:
        output = io.BytesIO()
        margin = 0.55 if scale < 1 else 0.7
        document = SimpleDocTemplate(output, pagesize=letter, leftMargin=margin * inch, rightMargin=margin * inch, topMargin=margin * inch, bottomMargin=margin * inch)
        styles = getSampleStyleSheet()
        name_style = ParagraphStyle("ResumeName", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=18 * scale, leading=21 * scale, alignment=1, spaceAfter=4 * scale)
        contact_style = ParagraphStyle("ResumeContact", parent=styles["Normal"], fontSize=9 * scale, leading=11 * scale, alignment=1, spaceAfter=10 * scale)
        heading_style = ParagraphStyle("ResumeHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=10.5 * scale, leading=13 * scale, spaceBefore=9 * scale, spaceAfter=4 * scale)
        body_style = ParagraphStyle("ResumeBody", parent=styles["Normal"], fontSize=10 * scale, leading=13 * scale, spaceAfter=3 * scale)
        story = []
        lines = [line.strip() for line in resume_text.splitlines() if line.strip()]
        for index, line in enumerate(lines):
            escaped = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
            if index == 0:
                story.append(Paragraph(escaped, name_style))
            elif index == 1:
                story.append(Paragraph(escaped, contact_style))
            elif is_section_heading(line):
                story.append(Paragraph(escaped.upper(), heading_style))
            elif line.startswith(("- ", "* ", "• ")):
                story.append(Paragraph(f"• {escaped[2:]}", body_style))
            else:
                story.append(Paragraph(escaped, body_style))
        story.append(Spacer(1, 1))
        document.build(story)
        return output.getvalue()


@dataclass(frozen=True)
class ResumeTemplate:
    font_name: str
    pdf_font_name: str
    body_size: float
    name_size: float
    heading_size: float
    line_height: float
    margin: float


RESUME_TEMPLATES = {
    "professional": ResumeTemplate("Aptos", "Helvetica", 10.5, 18, 11, 13, 0.7),
    "modern": ResumeTemplate("Arial", "Helvetica", 10.5, 20, 11, 13.5, 0.7),
    "classic": ResumeTemplate("Times New Roman", "Times-Roman", 10.5, 19, 11, 13.5, 0.75),
    "compact": ResumeTemplate("Aptos", "Helvetica", 9.5, 17, 10.5, 11.5, 0.58),
}


def resume_template(template_id: str) -> ResumeTemplate:
    return RESUME_TEMPLATES.get(template_id, RESUME_TEMPLATES["professional"])


def source_template(template: ResumeTemplate, profile) -> ResumeTemplate:
    if profile is None:
        return template
    value = profile.get if isinstance(profile, dict) else lambda key, default: getattr(profile, key, default)
    font_name = value("font_family", template.font_name)
    return ResumeTemplate(
        font_name if font_name in ALLOWED_FONT_FAMILIES else template.font_name,
        "Times-Roman" if font_name == "Times New Roman" else "Helvetica",
        float(value("body_size", template.body_size)),
        float(value("name_size", template.name_size)),
        float(value("heading_size", template.heading_size)),
        float(value("line_height", template.line_height)),
        template.margin,
    )


def _alignment(value: str) -> WD_ALIGN_PARAGRAPH:
    return {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(value, WD_ALIGN_PARAGRAPH.LEFT)


def _add_link(paragraph, run: TextRun) -> None:
    relationship_id = paragraph.part.relate_to(run.href, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    inline = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color"); color.set(qn("w:val"), "0563C1"); properties.append(color)
    underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); properties.append(underline)
    if run.bold: properties.append(OxmlElement("w:b"))
    if run.italic: properties.append(OxmlElement("w:i"))
    inline.append(properties)
    text = OxmlElement("w:t"); text.text = run.text; inline.append(text)
    hyperlink.append(inline)
    paragraph._p.append(hyperlink)


class RichResumeExportService(ResumeExportService):
    """Render editor HTML as one style system shared by DOCX and PDF exports."""

    def render_docx(self, resume_text: str, resume_html: str | None = None, target_page_count: int | None = None, template_id: str = "professional", style_profile=None) -> bytes:
        template = source_template(resume_template(template_id), style_profile) if template_id == "source" else resume_template(template_id)
        compact = target_page_count == 1 or template_id == "compact"
        document = Document()
        section = document.sections[0]
        margin = max(0.48, template.margin - (0.1 if compact else 0))
        section.top_margin = section.bottom_margin = Inches(margin)
        section.left_margin = section.right_margin = Inches(margin)
        normal = document.styles["Normal"]
        normal.font.name = template.font_name
        normal._element.rPr.rFonts.set(qn("w:eastAsia"), template.font_name)
        normal.font.size = Pt(template.body_size - (0.5 if compact else 0))
        normal.paragraph_format.line_spacing = Pt(template.line_height - (0.8 if compact else 0))
        normal.paragraph_format.space_after = Pt(2 if compact else 4)

        for index, block in enumerate(document_blocks(resume_text, resume_html)):
            text = "".join(run.text for run in block.runs)
            kind = "h1" if block.kind in {"h1", "h2"} else "h3" if block.kind == "h3" or (block.kind == "p" and is_section_heading(text)) else block.kind
            paragraph = document.add_paragraph(style="List Bullet" if kind == "li" else None)
            paragraph.alignment = _alignment("center" if kind == "h1" and index == 0 else block.alignment)
            paragraph.paragraph_format.space_before = Pt(7 if kind == "h3" else 0)
            paragraph.paragraph_format.space_after = Pt(3 if kind == "h3" else 2 if compact else 4)
            size = template.name_size if kind == "h1" else template.heading_size if kind == "h3" else template.body_size - (0.5 if compact else 0)
            for source_run in block.runs:
                if source_run.href:
                    _add_link(paragraph, source_run)
                    continue
                run = paragraph.add_run(source_run.text)
                run.font.name = template.font_name
                run.font.size = Pt(size)
                run.bold = source_run.bold or kind in {"h1", "h3"}
                run.italic = source_run.italic
                run.underline = source_run.underline
        output = io.BytesIO(); document.save(output)
        return output.getvalue()

    def render_source_docx(self, source_docx: bytes, resume_text: str, resume_html: str | None = None) -> bytes:
        """Patch body paragraphs in an uploaded DOCX while retaining its document setup and styles."""
        try:
            document = Document(io.BytesIO(source_docx))
        except Exception as error:
            raise HTTPException(status_code=422, detail="The original DOCX could not be opened.") from error

        blocks = document_blocks(resume_text, resume_html)
        paragraphs = [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]
        if not paragraphs:
            raise HTTPException(status_code=422, detail="The original DOCX has no editable body paragraphs.")
        templates = self._source_paragraph_templates(paragraphs)
        last_paragraph = paragraphs[-1]

        for index, block in enumerate(blocks):
            kind = self._source_block_kind(block)
            if index < len(paragraphs):
                paragraph = paragraphs[index]
                template = paragraph
            else:
                template = templates.get(kind) or last_paragraph
                paragraph = self._clone_paragraph_after(template, last_paragraph)
            self._replace_source_paragraph(paragraph, block, template)
            last_paragraph = paragraph

        for paragraph in paragraphs[len(blocks):]:
            element = paragraph._element
            element.getparent().remove(element)

        output = io.BytesIO()
        document.save(output)
        return output.getvalue()

    @staticmethod
    def _source_block_kind(block: ResumeBlock) -> str:
        text = "".join(run.text for run in block.runs)
        if block.kind in {"h1", "h2"}:
            return "title"
        if block.kind == "h3" or is_section_heading(text):
            return "heading"
        return "bullet" if block.kind == "li" else "body"

    @classmethod
    def _source_paragraph_templates(cls, paragraphs: list[DocxParagraph]) -> dict[str, DocxParagraph]:
        templates: dict[str, DocxParagraph] = {}
        for index, paragraph in enumerate(paragraphs):
            if index == 0:
                kind = "title"
            elif is_section_heading(paragraph.text):
                kind = "heading"
            elif paragraph._p.pPr is not None and paragraph._p.pPr.numPr is not None:
                kind = "bullet"
            else:
                kind = "body"
            templates.setdefault(kind, paragraph)
        templates.setdefault("body", paragraphs[0])
        return templates

    @staticmethod
    def _clone_paragraph_after(template: DocxParagraph, anchor: DocxParagraph) -> DocxParagraph:
        cloned = deepcopy(template._p)
        anchor._p.addnext(cloned)
        return DocxParagraph(cloned, anchor._parent)

    @staticmethod
    def _clear_paragraph_content(paragraph: DocxParagraph) -> None:
        for child in list(paragraph._p):
            if child.tag != qn("w:pPr"):
                paragraph._p.remove(child)

    @staticmethod
    def _copy_run_style(target, source) -> None:
        source_properties = source._r.rPr
        if source_properties is None:
            return
        target_properties = target._r.rPr
        if target_properties is not None:
            target._r.remove(target_properties)
        target._r.insert(0, deepcopy(source_properties))

    def _replace_source_paragraph(self, paragraph: DocxParagraph, block: ResumeBlock, template: DocxParagraph) -> None:
        source_run = next((run for run in template.runs if run.text), None)
        self._clear_paragraph_content(paragraph)
        if block.alignment in {"center", "right"}:
            paragraph.alignment = _alignment(block.alignment)
        for source in block.runs:
            if source.href:
                _add_link(paragraph, source)
                continue
            run = paragraph.add_run(source.text)
            if source_run is not None:
                self._copy_run_style(run, source_run)
            run.bold = bool(source.bold) or run.bold
            run.italic = bool(source.italic) or run.italic
            run.underline = bool(source.underline) or run.underline

    def render_pdf(self, resume_text: str, resume_html: str | None = None, target_page_count: int | None = None, template_id: str = "professional", style_profile=None) -> bytes:
        template = source_template(resume_template(template_id), style_profile) if template_id == "source" else resume_template(template_id)
        scales = (1.0, 0.95, 0.9, 0.85) if target_page_count == 1 else (1.0,)
        rendered = b""
        blocks = document_blocks(resume_text, resume_html)
        for scale in scales:
            rendered = self._render_rich_pdf(blocks, template, scale)
            if target_page_count is None or len(PdfReader(io.BytesIO(rendered)).pages) <= target_page_count:
                return rendered
        return rendered

    @staticmethod
    def _render_rich_pdf(blocks: list[ResumeBlock], template: ResumeTemplate, scale: float) -> bytes:
        output = io.BytesIO()
        margin = max(0.48, template.margin - (0.1 if scale < 1 else 0))
        document = SimpleDocTemplate(output, pagesize=letter, leftMargin=margin * inch, rightMargin=margin * inch, topMargin=margin * inch, bottomMargin=margin * inch)
        styles = getSampleStyleSheet()
        name_style = ParagraphStyle("RichResumeName", parent=styles["Title"], fontName=template.pdf_font_name, fontSize=template.name_size * scale, leading=(template.name_size + 3) * scale, alignment=1, spaceAfter=4 * scale)
        heading_style = ParagraphStyle("RichResumeHeading", parent=styles["Heading2"], fontName=template.pdf_font_name, fontSize=template.heading_size * scale, leading=(template.heading_size + 2) * scale, spaceBefore=9 * scale, spaceAfter=4 * scale)
        body_style = ParagraphStyle("RichResumeBody", parent=styles["Normal"], fontName=template.pdf_font_name, fontSize=template.body_size * scale, leading=template.line_height * scale, spaceAfter=3 * scale)
        story = []
        for index, block in enumerate(blocks):
            text = "".join(run.text for run in block.runs)
            kind = "h1" if block.kind in {"h1", "h2"} else "h3" if block.kind == "h3" or (block.kind == "p" and is_section_heading(text)) else block.kind
            fragments = []
            for run in block.runs:
                value = escape(run.text).replace("\n", "<br/>")
                if run.href: value = f'<a href="{escape(run.href, quote=True)}">{value}</a>'
                if run.underline: value = f"<u>{value}</u>"
                if run.italic: value = f"<i>{value}</i>"
                if run.bold or kind in {"h1", "h3"}: value = f"<b>{value}</b>"
                fragments.append(value)
            if kind == "li": fragments.insert(0, "&bull; ")
            style = name_style if kind == "h1" and index == 0 else heading_style if kind == "h3" else body_style
            story.append(Paragraph("".join(fragments), style))
        story.append(Spacer(1, 1)); document.build(story)
        return output.getvalue()
