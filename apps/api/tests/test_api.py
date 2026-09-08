import io

from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient

from rezzie.documents import ResumeExportService
from rezzie.main import app

client = TestClient(app)
LOCAL_IDENTITY = {"X-Rezzie-User-Id": "test-user"}


def test_health() -> None:
    assert client.get("/health").json() == {"status": "ok"}


def test_health_emits_security_headers() -> None:
    response = client.get("/health")
    assert response.headers["x-content-type-options"] == "nosniff"


def test_cors_preflight_allows_career_fact_updates() -> None:
    response = client.options(
        "/api/v1/career-records/record-1/facts/fact-1",
        headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "PATCH"},
    )
    assert response.status_code == 200
    assert "PATCH" in response.headers["access-control-allow-methods"]


def test_cors_preflight_allows_private_library_deletion() -> None:
    response = client.options(
        "/api/v1/resumes/resume-1",
        headers={"Origin": "http://localhost:5173", "Access-Control-Request-Method": "DELETE"},
    )
    assert response.status_code == 200
    assert "DELETE" in response.headers["access-control-allow-methods"]


def test_cors_preflight_is_not_rate_limited() -> None:
    response = client.options(
        "/api/v1/billing/checkout",
        headers={
            "Origin": "http://localhost:5173",
            "Access-Control-Request-Method": "POST",
            "Access-Control-Request-Headers": "authorization,content-type",
        },
    )
    assert response.status_code == 200
    assert response.headers["access-control-allow-origin"] == "http://localhost:5173"
    assert "POST" in response.headers["access-control-allow-methods"]


def test_text_import_validates_minimum_length() -> None:
    response = client.post("/api/v1/job-descriptions/text", headers=LOCAL_IDENTITY, json={"text": "short"})
    assert response.status_code == 422


def test_resume_file_imports_plain_text() -> None:
    response = client.post("/api/v1/resumes/file", headers=LOCAL_IDENTITY, files={"file": ("resume.txt", b"Experienced engineer with measurable delivery experience." * 2, "text/plain")})
    assert response.status_code == 200
    assert response.json()["source_type"] == "file"


def test_docx_resume_import_returns_a_safe_style_profile() -> None:
    source = Document()
    source.styles["Normal"].font.name = "Georgia"
    source.styles["Normal"].font.size = Pt(11)
    source.add_paragraph("Taylor Example")
    source.add_paragraph("taylor@example.com")
    source.add_paragraph("EXPERIENCE")
    role = source.add_paragraph(); role.add_run("Acme Corp | Engineer").bold = True
    source.add_paragraph("Delivered reliable systems and improved team workflows.")
    content = io.BytesIO(); source.save(content)

    response = client.post("/api/v1/resumes/file", headers=LOCAL_IDENTITY, files={"file": ("resume.docx", content.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})

    assert response.status_code == 200
    assert response.json()["style_profile"] == {
        "font_family": "Georgia", "body_size": 10.5, "line_height": 13.1,
        "name_size": 18.0, "heading_size": 10.5, "heading_uppercase": True,
        "emphasize_role_lines": True, "italic_metadata": False,
    }


def test_resume_pdf_import_reports_its_original_page_count() -> None:
    content = ResumeExportService().render_pdf("Taylor Example\ntaylor@example.com | Portland, OR\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems.")
    response = client.post("/api/v1/resumes/file", headers=LOCAL_IDENTITY, files={"file": ("resume.pdf", content, "application/pdf")})
    assert response.status_code == 200
    assert response.json()["page_count"] == 1


def test_resume_export_returns_an_editable_docx() -> None:
    response = client.post(
        "/api/v1/resumes/export",
        headers=LOCAL_IDENTITY,
        json={"resume_text": "Taylor Example\ntaylor@example.com | Portland, OR\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems."},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert response.headers["cache-control"] == "no-store"
    assert response.content[:2] == b"PK"


def test_resume_export_returns_a_pdf() -> None:
    response = client.post(
        "/api/v1/resumes/export/pdf",
        headers=LOCAL_IDENTITY,
        json={"resume_text": "Taylor Example\ntaylor@example.com | Portland, OR\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems."},
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("application/pdf")
    assert response.headers["cache-control"] == "no-store"
    assert response.content.startswith(b"%PDF")


def test_resume_export_accepts_a_one_page_target() -> None:
    response = client.post(
        "/api/v1/resumes/export/pdf",
        headers=LOCAL_IDENTITY,
        json={"resume_text": "Taylor Example\ntaylor@example.com | Portland, OR\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems.", "target_page_count": 1},
    )
    assert response.status_code == 200
    assert response.content.startswith(b"%PDF")


def test_billing_balance_requires_identity() -> None:
    assert client.get("/api/v1/billing/me").status_code == 401


def test_billing_balance_uses_local_development_identity() -> None:
    response = client.get("/api/v1/billing/me", headers={"X-Rezzie-User-Id": "local-user"})
    assert response.status_code == 200
    assert response.json()["purchased_credits"] == 0


def test_trusted_source_requires_ownership_attestation() -> None:
    response = client.post(
        "/api/v1/trusted-sources",
        headers=LOCAL_IDENTITY,
        json={"label": "Portfolio", "url": "https://example.com", "ownership_attested": False},
    )
    assert response.status_code == 422
    assert "Confirm that you own" in response.json()["detail"]


def test_trusted_source_requires_subscription_before_fetching() -> None:
    response = client.post(
        "/api/v1/trusted-sources",
        headers=LOCAL_IDENTITY,
        json={"label": "Portfolio", "url": "https://example.com", "ownership_attested": True},
    )
    assert response.status_code == 403


def test_subscription_route_requires_server_model_key() -> None:
    response = client.post("/api/v1/tailor", json={"resume_text": "a" * 50, "job_description": "b" * 50, "credential_mode": "subscription"})
    assert response.status_code == 503
