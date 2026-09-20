import io

from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader

from rezzie.documents import (
    ResumeExportService,
    RichResumeExportService,
    docx_style_profile,
    editor_html_matches_resume_text,
    is_section_heading,
    paragraph_text,
)


def test_document_extraction_keeps_paragraph_and_table_boundaries() -> None:
    source = Document()
    source.add_paragraph("Taylor Example")
    source.add_paragraph("EXPERIENCE")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Acme Corp"
    table.cell(0, 1).text = "Engineer"

    assert (
        paragraph_text(source) == "Taylor Example\n\nEXPERIENCE\n\nAcme Corp | Engineer"
    )


def test_exported_docx_has_resume_structure() -> None:
    content = ResumeExportService().render_docx(
        "Taylor Example\ntaylor@example.com | Portland, OR\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems."
    )
    rendered = Document(io.BytesIO(content))

    assert rendered.paragraphs[0].text == "Taylor Example"
    assert rendered.paragraphs[0].runs[0].bold
    assert "EXPERIENCE" in [paragraph.text for paragraph in rendered.paragraphs]
    assert any(
        paragraph.text == "Delivered reliable systems."
        for paragraph in rendered.paragraphs
    )
    assert is_section_heading("Experience:")
    assert is_section_heading("Professional Experience")


def test_rich_editor_formatting_survives_docx_and_pdf_exports() -> None:
    source_html = "<h1>Taylor Example</h1><p><strong>taylor@example.com</strong> · <em>Portland, OR</em></p><h3>EXPERIENCE</h3><ul><li><u>Delivered</u> reliable systems.</li></ul>"
    service = RichResumeExportService()

    rendered = Document(
        io.BytesIO(
            service.render_docx("Taylor Example", source_html, template_id="classic")
        )
    )
    contact_runs = rendered.paragraphs[1].runs

    assert rendered.paragraphs[0].runs[0].font.name == "Times New Roman"
    assert contact_runs[0].bold
    assert any(run.italic for run in contact_runs)
    assert rendered.paragraphs[-1].runs[0].underline
    assert service.render_pdf(
        "Taylor Example", source_html, template_id="modern"
    ).startswith(b"%PDF")


def test_rich_exports_preserve_clickable_links() -> None:
    source_html = '<h1>Taylor Example</h1><p><a href="https://portfolio.example.com">portfolio.example.com</a></p>'
    rendered = Document(
        io.BytesIO(RichResumeExportService().render_docx("Taylor Example", source_html))
    )

    assert any(
        str(relationship.target_ref) == "https://portfolio.example.com"
        for relationship in rendered.part.rels.values()
    )
    assert (
        RichResumeExportService()
        .render_pdf("Taylor Example", source_html)
        .startswith(b"%PDF")
    )


def test_rich_exports_linkify_plain_text_urls_when_editor_html_is_unavailable() -> None:
    resume_text = "Taylor Example\nwww.portfolio.example.com | mailto:taylor@example.com\n\nEXPERIENCE\n- Built https://github.com/example/project."
    service = RichResumeExportService()
    rendered = Document(io.BytesIO(service.render_docx(resume_text)))
    targets = {
        str(relationship.target_ref) for relationship in rendered.part.rels.values()
    }

    assert "https://www.portfolio.example.com" in targets
    assert "mailto:taylor@example.com" in targets
    assert "https://github.com/example/project" in targets
    pdf = PdfReader(io.BytesIO(service.render_pdf(resume_text)))
    targets_in_pdf = {
        str(annotation.get_object()["/A"]["/URI"])
        for page in pdf.pages
        for annotation in (page.get("/Annots") or [])
        if annotation.get_object().get("/A", {}).get("/URI")
    }
    assert {
        "https://www.portfolio.example.com",
        "mailto:taylor@example.com",
        "https://github.com/example/project",
    } <= targets_in_pdf


def test_editor_html_match_accepts_browser_list_text_without_markers() -> None:
    editor_html = "<h1>Taylor Example</h1><h3>SKILLS</h3><ul><li><strong>TypeScript</strong> and React</li><li>System design</li></ul>"

    assert editor_html_matches_resume_text(
        editor_html,
        "Taylor Example\nSKILLS\nTypeScript and React\nSystem design",
    )


def test_editor_html_match_accepts_a_browser_flattened_block_snapshot() -> None:
    editor_html = "<h1>Taylor Example</h1><h3>SUMMARY</h3><p>Builds reliable systems.</p><ul><li>TypeScript and React</li></ul>"

    assert editor_html_matches_resume_text(
        editor_html,
        "Taylor ExampleSUMMARYBuilds reliable systems.TypeScript and React",
    )


def test_source_docx_export_preserves_document_setup_and_patches_its_paragraphs() -> (
    None
):
    source = Document()
    source.sections[0].left_margin = Inches(1.1)
    source.styles["Normal"].font.name = "Georgia"
    name = source.add_paragraph("Taylor Example")
    name.runs[0].font.size = Pt(20)
    source.add_paragraph("taylor@example.com")
    source.add_paragraph("SUMMARY")
    summary = source.add_paragraph("Original summary.")
    summary.runs[0].font.size = Pt(10)
    source.add_paragraph("EXPERIENCE")
    source.add_paragraph("Acme Corp | Engineer")
    source.add_paragraph("Delivered reliable systems.", style="List Bullet")
    source_bytes = io.BytesIO()
    source.save(source_bytes)

    output = RichResumeExportService().render_source_docx(
        source_bytes.getvalue(),
        "Taylor Example\ntaylor@example.com\n\nSUMMARY\nGrounded systems engineer.\n\nEXPERIENCE\nAcme Corp | Engineer\n- Delivered reliable systems and improved deployment speed.",
    )
    rendered = Document(io.BytesIO(output))

    assert rendered.sections[0].left_margin == Inches(1.1)
    assert rendered.styles["Normal"].font.name == "Georgia"
    assert "Grounded systems engineer." in [
        paragraph.text for paragraph in rendered.paragraphs
    ]
    assert any(
        "improved deployment speed" in paragraph.text
        for paragraph in rendered.paragraphs
    )
    rendered_summary = next(
        paragraph
        for paragraph in rendered.paragraphs
        if paragraph.text == "Grounded systems engineer."
    )
    assert rendered_summary.runs[0].font.size == Pt(10)


def test_source_docx_export_preserves_source_link_targets_for_retained_labels() -> None:
    source = RichResumeExportService().render_docx(
        "Taylor Example",
        '<h1>Taylor Example</h1><p>Portfolio: <a href="https://portfolio.example.com/work">Selected work and grounded engineering case studies</a></p>',
    )

    output = RichResumeExportService().render_source_docx(
        source,
        "Taylor Example\nPortfolio: Selected work and grounded engineering case studies",
    )
    rendered = Document(io.BytesIO(output))

    assert any(
        str(relationship.target_ref) == "https://portfolio.example.com/work"
        for relationship in rendered.part.rels.values()
    )


def test_docx_style_profile_extracts_portable_hierarchy() -> None:
    source = Document()
    source.styles["Normal"].font.name = "Georgia"
    source.styles["Normal"].font.size = Pt(11)
    name = source.add_paragraph()
    name.add_run("Taylor Example").font.size = Pt(20)
    source.add_paragraph("taylor@example.com")
    heading = source.add_paragraph()
    heading_run = heading.add_run("EXPERIENCE")
    heading_run.bold = True
    heading_run.font.size = Pt(12)
    role = source.add_paragraph()
    role.add_run("Acme Corp | Engineer").bold = True
    dates = source.add_paragraph()
    dates_run = dates.add_run("2020 - 2024")
    dates_run.italic = True

    profile = docx_style_profile(source)

    assert profile["font_family"] == "Georgia"
    assert profile["name_size"] == 20
    assert profile["heading_size"] == 12
    assert profile["emphasize_role_lines"] is True
    assert profile["italic_metadata"] is True
